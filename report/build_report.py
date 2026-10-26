"""Build HashLab's ten-page black-and-white academic report."""

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from experiments import run_experiment
from figures import create_figures

OUT=ROOT/'output/report'
TMP=ROOT/'tmp/report/compact'
OUT.mkdir(parents=True,exist_ok=True)
TMP.mkdir(parents=True,exist_ok=True)
runs={}
for distribution in ['Random keys','Same initial bucket']:
    for fill in [.9,1.0]:
        rows,summary,keys=run_experiment(101,distribution,42,fill)
        runs[f'{distribution}_{fill}']={'rows':rows,'summary':summary,'keys':keys}
(TMP/'measurements.json').write_text(json.dumps(runs,indent=2),encoding='utf-8')
figures=create_figures(TMP/'figures',runs)

doc=Document()
s=doc.sections[0]
s.page_width,s.page_height=Inches(8.2677),Inches(11.6929)
s.top_margin,s.bottom_margin=Inches(.78),Inches(.72)
s.left_margin,s.right_margin=Inches(.85),Inches(.85)
s.footer_distance=Inches(.32)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name]
    st.font.name='Times New Roman'
    st.font.color.rgb=RGBColor(0,0,0)
    fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+attribute),None)
    for attribute in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+attribute),'Times New Roman')
normal=doc.styles['Normal']
normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.widow_control=True
for name,size in [('Title',18),('Heading 1',16),('Heading 2',12.5)]:
    st=doc.styles[name]
    st.font.size=Pt(size)
    st.font.bold=True
    st.paragraph_format.space_before=Pt(9)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
caption=doc.styles['Caption']
caption.font.size=Pt(10.5)
caption.font.bold=False
caption.font.italic=True
caption.paragraph_format.space_after=Pt(8)
caption.paragraph_format.space_before=Pt(3)
footer=s.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

page_count=0
figure_count=0
table_count=0


def page(title,style='Heading 1'):
    global page_count
    page_count+=1
    para=doc.add_paragraph(title,style)
    para.paragraph_format.page_break_before=page_count>1
    return para


def sub(text):
    return doc.add_paragraph(text,'Heading 2')


def p(text):
    para=doc.add_paragraph(text)
    para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return para


def bullets(items):
    for text in items:
        para=doc.add_paragraph(text,'List Bullet')
        para.paragraph_format.space_after=Pt(4)


def code(text):
    para=doc.add_paragraph()
    para.paragraph_format.keep_together=True
    para.paragraph_format.line_spacing=1
    para.paragraph_format.space_after=Pt(8)
    run=para.add_run(text);run.font.name='Consolas';run.font.size=Pt(10)


def equation(text,fraction=None):
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMath')
    def run(parent,value):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);parent.append(r)
    run(math,text)
    if fraction:
        f=OxmlElement('m:f')
        for tag,value in zip(['m:num','m:den'],fraction):
            child=OxmlElement(tag);run(child,value);f.append(child)
        math.append(f)
    para._p.append(math)
    para.paragraph_format.space_after=Pt(7)


def figure(key,caption_text,width):
    global figure_count
    figure_count+=1
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    para.paragraph_format.space_after=Pt(1)
    drawing=para.add_run().add_picture(str(figures[key]),width=Inches(width))
    drawing._inline.docPr.set('descr',caption_text)
    para=doc.add_paragraph(f'Figure {figure_count}. {caption_text}','Caption')
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER


def table(title,headers,rows,widths):
    global table_count
    table_count+=1
    cap=doc.add_paragraph(f'Table {table_count}. {title}','Caption')
    cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers))
    t.autofit=False
    numeric_columns={j for j in range(len(headers)) if all(str(row[j]).replace('.', '', 1).isdigit() for row in rows)}
    for c,w,value in zip(t.rows[0].cells,widths,headers):
        c.width=Inches(w);c.text=value
    for values in rows:
        for c,w,value in zip(t.add_row().cells,widths,values):
            c.width=Inches(w);c.text=str(value)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        if i==0:
            trPr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge)
                for key,value in [('val','single'),('sz','4'),('color','808080')]:
                    node.set(qn('w:'+key),value)
                borders.append(node)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),'65');node.set(qn('w:type'),'dxa');margins.append(node)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(1)
                para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.line_spacing=1
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER if j in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(11)
                    run.font.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


# Page 1 includes the abstract and contents. There is no cover page.
page('Abstract','Title')
p('HashLab is a Python and Streamlit application that demonstrates hashing through separate chaining, linear probing and double hashing. Users can insert, search and delete integer keys, inspect the table, and follow the most recent operation. A library mode uses book IDs to retrieve titles and shelf locations through the same custom hash-table engine.')
p('The project studies how load factor and input distribution affect collision handling. Identical random and deliberately colliding datasets are used to compare insertion and search costs. The measurements count key comparisons and array visits rather than machine-dependent execution time. Results show that long chains and probe clusters increase work, while double hashing spreads the selected colliding keys more effectively. Both probing methods can inspect the entire table during an unsuccessful search at full capacity.')
p('Thirteen automated tests passed, covering table operations, book records, duplicate IDs, deletion, full tables and interface workflows. The system is a demonstration application with session-based records; it is not a permanent library catalogue.')
para=doc.add_paragraph('Contents','Title')
para.paragraph_format.space_before=Pt(13)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u')
doc.add_paragraph()._p.append(field)

page('1 Introduction and Requirements')
sub('1.1 Background and problem statement')
p('Hashing provides a way to locate a record from its key without scanning every record. A hash function converts the key into an array index. Since the number of possible keys is larger than the number of array positions, different keys can reach the same index. The table must resolve this collision without losing or confusing the records.')
p('The project makes this behaviour visible and measurable. It shows where an operation begins, which positions it visits and how many stored keys it compares. The library example gives these operations a practical purpose. Book IDs 12 and 23 both start at index 1 in a table of size 11, but a search must still return the correct title and shelf for the requested ID.')
sub('1.2 Objectives')
bullets(['Implement insertion, search and deletion using three collision strategies.','Display the table and the steps of the most recent operation.','Compare strategies using the same keys, order, capacity and search queries.','Demonstrate record lookup through a small library example.'])
sub('1.3 Scope and requirements')
p('The Simulator supports integer keys and library books. Users can choose a prime capacity, load five colliding entries, clear the table and inspect a trace. The comparison tab generates random or same-bucket input and exports measurements as CSV. Borrowing, fines, accounts, editing records and permanent storage are outside the current scope.')
p('The application requires Python 3.10 or newer, Streamlit, pandas and Altair. The version ranges are recorded in requirements.txt. It runs locally in a browser and needs no database, API key or external service after dependency installation.')
p('Library IDs must be positive integers. Adding a book requires a non-blank title and shelf. Spaces at the start and end are removed; the fields are limited to 120 and 40 characters respectively. A repeated ID keeps its existing record. Changing mode, strategy or capacity clears the simulator, as does loading a replacement sample.')

page('2 Hashing and Collision Resolution')
sub('2.1 Hash function and load factor')
p('HashLab uses the division method, where m is the number of slots or buckets:')
equation('h(k) = k mod m')
p('The index lies between 0 and m - 1. It is a starting location, not a unique identity. For example, 12, 23 and 34 all have remainder 1 when divided by 11.')
figure('collision','Different book IDs sharing an initial index',6.2)
p('If n keys are stored, the load factor is:')
equation('\u03b1 = ',('n','m'))
p('Five entries in 11 buckets give a load factor of about 0.45. Open addressing has one live key per slot, so its load cannot exceed 1. Chaining can exceed 1 because each bucket holds a linked list.')
sub('2.2 Implemented strategies')
table('Collision strategies in HashLab',['Strategy','Storage','Next candidate'],[
    ['Separate chaining','Linked nodes in a bucket','Next node'],
    ['Linear probing','One key per array slot','Next slot'],
    ['Double hashing','One key per array slot','Key-dependent step'],
],[1.65,2.55,2.36])
p('Linear probing uses step 1. Double hashing uses 1 + (k mod (m - 1)). All capacities are prime, so every possible nonzero step is relatively prime to m and can visit all slots. This complete-cycle condition is important for correct probing [1]. Standard algorithm analysis describes good expected lookup cost under suitable hashing and load assumptions [1].')

page('3 System Design and Workflow')
p('The system has three main modules. app.py handles the Streamlit controls, validation and session state. hashing.py implements the custom tables and returns results and traces. experiments.py calls the same engine with controlled datasets. styles.css handles display styling.')
figure('workflow','Complete workflow of the simulator and comparison module',6.25)
p('The table, previous snapshot, latest result, collision total and trace position are held in session state [2]. They last for the current session. Comparison results are stored separately from the simulator. A Result contains the status, visits, comparisons, insertion collisions, trace steps and, when available, the matching book.')
p('The operation selector is outside the input form so that it can show the correct fields immediately. Submitted values are processed before the table and results are drawn. Forms group values into one submission [2], and no extra forced rerun is needed.')

page('4 Algorithm Implementation')
sub('4.1 Separate chaining')
p('Each bucket stores the head of a linked list. A node holds a key, an optional Book record and the next reference. Insertion scans for a duplicate before appending a node. Search stops at a full-key match or the end of the list. Deletion unlinks only the matching node, including the head when necessary.')
sub('4.2 Linear probing and double hashing')
p('Both probing methods use the following sequence, with i starting at zero:')
equation('p(i, k) = (h(k) + i * step(k)) mod m')
code('PROBE_INSERT(key, book)\n    first_deleted = NONE\n    for index in at most m probe positions:\n        if slot is DELETED:\n            remember first deleted index; continue\n        if slot is EMPTY:\n            store in first_deleted or current index; return\n        if slot contains key: return ALREADY_PRESENT\n        count the unequal live key as a collision\n    if first_deleted exists: store there\n    otherwise: return TABLE_FULL')
p('A deleted slot is remembered, but does not end the duplicate check. Search also continues past deleted markers. A never-used empty slot ends a failed search. Every probe loop is bounded by m positions, so a full table cannot cause an infinite loop.')
figure('tombstone','Deletion preserves the path to later keys',6.2)
p('After removing 12, a search for 34 must pass the marker at slot 1 and compare the later keys. Treating slot 1 as empty would stop the search too early. Deletion also clears the book record at that slot. Matching always checks the full ID rather than only its hash index.')

page('5 Library Book Lookup')
sub('5.1 Records and operations')
p('Library mode supports Add book, Find book and Remove book. The ID is the integer key; the record contains a title and shelf location. Chaining stores the record in its linked node. Probing stores the record at the same array index as the key. Book retrieval therefore uses the custom table, not a separate Python dictionary.')
table('Collision sample with capacity 11',['Book ID','Title','Shelf'],[
    [12,'Algorithms','A2'],[23,'Python Basics','B3'],[34,'Data Structures','C4'],
    [45,'Computer Networks','D5'],[56,'Database Systems','E6'],
],[1,4.46,1.1])
p('All five sample IDs start at index 1. Chaining stores them in one list. Linear probing stores them in slots 1 to 5. Double hashing places them at indices 1, 5, 6, 7 and 8 respectively. These are demonstration records, not a real library dataset.')
figure('book_lookup','Finding the book with ID 23 using separate chaining',5.7)
sub('5.2 Duplicate and deletion behaviour')
p('Finding ID 23 returns Python Basics on shelf B3. With the sample loaded, this takes two comparisons in each strategy. Chaining visits one bucket; either probing method visits two slots. After removing ID 12, ID 23 remains searchable. Chaining needs one comparison; probing still visits two slots but makes only one live-key comparison because the first slot is deleted.')
p('Adding an existing ID reports Already present and keeps its title and shelf unchanged. Removing an absent ID reports Not found and does not change the entry count. Records are session-only, and switching the table configuration starts a new table.')

page('6 Complexity and Implementation Details')
sub('6.1 Correctness and termination')
p('In chaining, every key remains in the bucket selected by its initial hash. Insertion adds a node only after checking for an equal key; deletion changes links without disconnecting the remaining records. In probing, insertion and search use the same sequence. A valid step visits all slots, and deleted markers preserve the sequence after removal.')
p('A never-used empty position safely ends search because insertion would have stopped at that position before placing the same key farther along. A full cycle ends after m visits. Equal starting indices are never enough to identify a book: only a full-key match returns its associated record.')
sub('6.2 Time and space requirements')
table('Complexity of the implemented tables',['Measure','Separate chaining','Open addressing'],[
    ['Expected search','O(1 + alpha), under uniform hashing','Expected O(1) with suitable hashing and controlled load'],
    ['Worst operation','O(n) node checks','O(m) slot visits'],
    ['Table storage','O(m + n), plus record data','O(m), plus record data'],
    ['Longest trace','O(n)','O(m)'],
],[1.35,2.55,2.66])
p('Here n is the number of keys and m is capacity. If n and m are of the same order, a full probe cycle is linear in the collection size. Chaining may have n greater than m. This implementation scans a chain before appending, so insertion into a long chain is not unconditionally constant time.')
sub('6.3 Counting rules and display cost')
p('A bucket or slot visit counts an inspected array position. A key comparison counts equality testing against a live stored key. Chaining visits one array bucket and then compares linked nodes. An insertion collision counts an unequal live key encountered during an insertion attempt. Empty positions and deleted markers cost visits but no live-key comparison.')
p('The interface displays the state before an operation at earlier trace steps and the completed state at the final step. It does not animate every intermediate mutation. Snapshot creation and drawing have their own cost and are not included in the engine counters. Experiments disable trace collection to keep the comparison focused on table operations.')

page('7 Experimental Results')
sub('7.1 Method')
p('The reported runs use capacity 101 and seed 42. Each strategy receives identical unique keys, order and queries. At 90% target fill, 90 keys are inserted, giving actual load 90/101 = 0.8911. Random keys and absent queries come from disjoint integer ranges using random.sample. Colliding keys follow 3 + 101i; absent queries continue beyond the inserted range.')
figure('results','Average insertion comparisons as the load factor increases',6.5)
p('Up to ten checkpoints are measured. Insertion averages cover all keys inserted so far. Successful searches cover every present key; unsuccessful searches use the same fixed sample of 100 absent keys. Table 4 reports average live-key comparisons at the final checkpoint.')
rows=[]
for distribution,label in [('Random keys','Random'),('Same initial bucket','Colliding')]:
    for r in runs[distribution+'_0.9']['summary']:
        rows.append([label,r['Strategy'],f"{r['Insertion comparisons']:.2f}",f"{r['Successful search comparisons']:.2f}",f"{r['Unsuccessful search comparisons']:.2f}"])
table('Average comparisons at 90 keys and capacity 101',['Input','Strategy','Insert','Found','Absent'],rows,[.85,2.11,1.2,1.2,1.2])
p('With colliding input, chaining and linear probing each make 4,005 insertion comparisons, or 44.50 per key. Double hashing makes 89, averaging 0.99. Its secondary steps spread these particular keys. Keys sharing both hash components can still follow the same probe path, so this is not a worst-case guarantee.')

page('8 Testing and Discussion')
sub('8.1 Automated validation')
p('Thirteen unittest and Streamlit AppTest tests passed: eight cover the engine and experiments, and five cover interface workflows. The main checks are:')
bullets(['Colliding keys remain searchable, and duplicates do not increase the count.','Deletion preserves other records; full probing stops and deleted slots can be reused.','Double-hash sequences cover every slot for 200 tested keys at five prime capacities.','Eight hundred random operations per strategy agree with a reference set.','The same experiment settings reproduce the same measurements.','Book details survive collisions, duplicates, deletion and slot reuse.','The interface validates input, loads samples, runs comparisons and resets correctly.'])
p('The reference set is used only as a test oracle. It does not replace the custom table in the application. The library tests confirm that a full-table insertion failure does not create a hidden record and that an existing ID retains its original title.')
sub('8.2 Interpretation of the results')
p('The random workload produces short average chains, while probing becomes more expensive as occupancy rises. In the deliberately colliding workload, chaining forms one long list and linear probing forms one cluster. Their equal comparison counts do not imply equal execution time or memory behaviour.')
p('A separate random run stores 101 keys at capacity 101. Both probing strategies then make 101 comparisons for every absent query because no empty stopping slot remains. Chaining averages 1.10 comparisons for the same absent-query sample. Its capacity describes the number of buckets, not a fixed limit on nodes.')
p('The experiments show that both occupancy and key distribution matter. They use one seed and constructed datasets, and do not measure elapsed time, memory consumption or usability with library staff. Passing tests supports correctness in the tested cases; it does not establish production performance.')
sub('8.3 Reproducing the application and tests')
code('python -m pip install -r requirements.txt\npython -m streamlit run app.py\npython -m unittest discover -s tests -v')

page('9 Conclusion and References')
sub('9.1 Conclusion')
p('HashLab connects collision-resolution algorithms with visible table operations and a small library lookup. It preserves the correct book record when different IDs share an index. Chaining uses linked nodes; probing uses bounded array sequences and deleted markers. The same engine supplies the simulator and comparison module.')
p('The measurements show increasing work in long chains and occupied probe regions. Double hashing spreads the chosen colliding sequence more effectively, but it still scans the full table for an absent key at full capacity. Good expected performance therefore depends on suitable hashing, the input distribution and controlled load.')
sub('9.2 Limitations and future work')
p('Records last only for the current session. The application has numeric-ID lookup but no title search, editing, lending or permanent catalogue. The table does not resize or clean accumulated deleted markers. Useful extensions would be a small persistence layer, optional rehashing, and comparisons across more seeds and keys that share both hash components. These should preserve the project\'s focus on explaining the data structure.')
sub('References')
refs=[
    '[1] Cormen et al. Introduction to Algorithms.',
    '[2] Streamlit documentation. https://docs.streamlit.io/',
]
for ref in refs:
    para=doc.add_paragraph(ref)
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.left_indent=Inches(.18)
    para.paragraph_format.first_line_indent=Inches(-.18)
    para.paragraph_format.space_after=Pt(5)
    for run in para.runs:
        run.font.size=Pt(10.5)
para=p('Figures and measurements were generated from the HashLab implementation.')
for run in para.runs:
    run.font.size=Pt(10.5)

for style in doc.styles:
    if style.type==1:
        for border in list(style.element.findall('.//'+qn('w:pBdr'))):
            border.getparent().remove(border)
for para in doc.paragraphs:
    if para.style.name in ['Title','Heading 1','Heading 2']:
        for run in para.runs:
            run.font.name='Times New Roman';run.font.bold=True;run.font.color.rgb=RGBColor(0,0,0)
        borders=OxmlElement('w:pBdr')
        for edge in ['top','bottom','left','right']:
            node=OxmlElement('w:'+edge);node.set(qn('w:val'),'nil');borders.append(node)
        para._p.get_or_add_pPr().append(borders)
doc.core_properties.title='HashLab Hash Collision Simulator and Library Book Lookup'
doc.core_properties.subject='Design and Analysis of Algorithms project report'
doc.core_properties.author=''
doc.save(OUT/'HashLab_Project_Report.docx')
print(f'Built {page_count} planned pages, {figure_count} monochrome figures and {table_count} simple tables.')
